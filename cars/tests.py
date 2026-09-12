from io import BytesIO
from types import SimpleNamespace
from unittest.mock import patch

from django.contrib.auth.models import User
from django.core.cache import cache, caches
from django.db import IntegrityError
from django.test import Client, TestCase
from django.test.utils import override_settings
from django.utils import timezone

from cars.models import CarSpecification, Dealer, PromoCode, SiteSettings, Sponsor, SITE_SETTINGS_CACHE_KEY
from cars.services.deepseek_service import _provider_chain
from cars.views import _client_ip

# الاختبارات تعمل في عملية واحدة، لذا نستبدل التخزين "المشترك" بذاكرة محلية
# لنفس المنطق (العدّاد والحد) دون الحاجة لجدول قاعدة بيانات في قاعدة الاختبار.
_LOCMEM_CACHES = {
    'default': {'BACKEND': 'django.core.cache.backends.locmem.LocMemCache', 'LOCATION': 't-default'},
    'shared': {'BACKEND': 'django.core.cache.backends.locmem.LocMemCache', 'LOCATION': 't-shared'},
}


def _make_sponsor(slug='testco', prefix='TEST'):
    sp = Sponsor.objects.create(name='شركة تجريبية', slug=slug, code_prefix=prefix, discount=10)
    sp.set_password('secret123')
    sp.save()
    return sp


def _make_car(car_id=1, brand_ar='تويوتا', model_ar='كورولا'):
    return CarSpecification.objects.create(
        id=car_id,
        brand_en='Toyota',
        brand_ar=brand_ar,
        model_en='Corolla',
        model_ar=model_ar,
        year=2020,
        spec='خليجي',
        engine='1.8',
        oil_visc='5W-30',
        fuel='بنزين',
        octane=91,
        tire_size='205/55R16',
        oil_capacity='4.2L',
    )


class SiteSettingsCacheTests(TestCase):
    """قراءة الإعدادات من الذاكرة المؤقتة وإبطالها عند الحفظ."""

    def test_load_is_cached(self):
        cache.delete(SITE_SETTINGS_CACHE_KEY)
        SiteSettings.load()
        # القراءة الثانية يجب أن تأتي من الذاكرة المؤقتة دون أي استعلام قاعدة بيانات
        with self.assertNumQueries(0):
            second = SiteSettings.load()
        self.assertEqual(second.pk, 1)
        self.assertIsNotNone(cache.get(SITE_SETTINGS_CACHE_KEY))
        cache.delete(SITE_SETTINGS_CACHE_KEY)

    def test_save_invalidates_cache(self):
        cache.delete(SITE_SETTINGS_CACHE_KEY)
        obj = SiteSettings.load()
        self.assertIsNotNone(cache.get(SITE_SETTINGS_CACHE_KEY))
        obj.site_name = 'نسخة مختبرية'
        obj.save()
        self.assertIsNone(cache.get(SITE_SETTINGS_CACHE_KEY))


class PromoCodeConstraintTests(TestCase):
    def test_code_is_unique(self):
        sponsor = Sponsor.objects.create(name='زيت الحسام', slug='hisam', code_prefix='HISAM')
        PromoCode.objects.create(code='HISAM-1111', sponsor=sponsor)
        with self.assertRaises(IntegrityError):
            PromoCode.objects.create(code='HISAM-1111', sponsor=sponsor)


class ClientIpTests(TestCase):
    """العنوان الحقيقي يؤخذ من آخر X-Forwarded-For (nginx يضيفه آخراً)."""

    def test_takes_last_xff_not_first(self):
        req = SimpleNamespace(META={'HTTP_X_FORWARDED_FOR': '1.2.3.4, 203.0.113.9'})
        self.assertEqual(_client_ip(req), '203.0.113.9')

    def test_falls_back_to_remote_addr(self):
        req = SimpleNamespace(META={'REMOTE_ADDR': '127.0.0.1'})
        self.assertEqual(_client_ip(req), '127.0.0.1')


@override_settings(CACHES=_LOCMEM_CACHES)
class PromoGenerationTests(TestCase):
    def setUp(self):
        self.sp = _make_sponsor()
        caches['shared'].delete('codegen:203.0.113.5')
        self.client = Client(REMOTE_ADDR='203.0.113.5')

    def test_generate_success(self):
        r = self.client.post('/api/promo-code/generate/', {'sponsor': self.sp.slug})
        self.assertEqual(r.status_code, 200)
        data = r.json()
        self.assertTrue(data['success'])
        self.assertTrue(PromoCode.objects.filter(code=data['code'], sponsor=self.sp).exists())

    def test_generate_unknown_sponsor_returns_404(self):
        r = self.client.post('/api/promo-code/generate/', {'sponsor': 'nope'})
        self.assertEqual(r.status_code, 404)

    def test_generate_returns_same_code_for_same_ip(self):
        r1 = self.client.post('/api/promo-code/generate/', {'sponsor': self.sp.slug})
        r2 = self.client.post('/api/promo-code/generate/', {'sponsor': self.sp.slug})
        self.assertEqual(r1.status_code, 200)
        self.assertEqual(r2.status_code, 200)
        self.assertEqual(r1.json()['code'], r2.json()['code'])
        self.assertEqual(r2.json()['status'], 'existing')
        self.assertEqual(PromoCode.objects.filter(sponsor=self.sp).count(), 1)

    def test_rate_limit_per_ip(self):
        caches['shared'].set('codegen:203.0.113.5', 60, 3600)
        r = self.client.post('/api/promo-code/generate/', {'sponsor': self.sp.slug})
        self.assertEqual(r.status_code, 429)
        self.assertFalse(r.json()['success'])

    def test_rate_counter_increments_after_success(self):
        self.client.post('/api/promo-code/generate/', {'sponsor': self.sp.slug})
        self.assertEqual(caches['shared'].get('codegen:203.0.113.5', 0), 1)


@override_settings(CACHES=_LOCMEM_CACHES)
class ServicesLoginTests(TestCase):
    def setUp(self):
        self.sp = _make_sponsor()
        self.ip = '203.0.113.7'
        caches['shared'].delete(f'sfail:{self.ip}')

    def test_lockout_after_five_failures(self):
        c = Client(REMOTE_ADDR=self.ip)
        for _ in range(5):
            c.post('/services/', {'identifier': 'wrong', 'password': 'x'})
        r = c.post('/services/', {'identifier': 'wrong', 'password': 'x'})
        self.assertContains(r, 'محاولات كثيرة')

    def test_unsuccessful_attempt_counts_up(self):
        Client(REMOTE_ADDR=self.ip).post('/services/', {'identifier': 'x', 'password': 'x'})
        self.assertEqual(caches['shared'].get(f'sfail:{self.ip}', 0), 1)

    def test_success_clears_counter(self):
        caches['shared'].set(f'sfail:{self.ip}', 4, 300)
        c = Client(REMOTE_ADDR=self.ip)
        r = c.post('/services/', {'identifier': self.sp.slug, 'password': 'secret123'})
        self.assertEqual(r.status_code, 302)
        self.assertIsNone(caches['shared'].get(f'sfail:{self.ip}'))


class MixCalculatorTests(TestCase):
    def test_valid_calculation_message(self):
        r = self.client.post('/mix/', {
            'octane_target': '95', 'octane1': '91', 'octane2': '98', 'tank_capacity': '50',
        })
        self.assertEqual(r.status_code, 200)
        self.assertContains(r, 'تم حساب الخلطة بنجاح')

    def test_invalid_tank_shows_error(self):
        r = self.client.post('/mix/', {
            'octane_target': '95', 'octane1': '91', 'octane2': '98', 'tank_capacity': '0',
        })
        self.assertContains(r, 'سعة الخزان')

    def test_non_numeric_values_show_error(self):
        r = self.client.post('/mix/', {
            'octane_target': 'abc', 'octane1': '91', 'octane2': '98', 'tank_capacity': '50',
        })
        self.assertContains(r, 'أرقام صحيحة')


class ReportViewTests(TestCase):
    """تقرير الأكواد: يحمّل لحد أقصى ويقبل معاملات تالفة دون 500."""

    def setUp(self):
        self.staff = User.objects.create_superuser('boss', 'boss@example.com', 'pw')
        self.client.force_login(self.staff)
        self.sp = _make_sponsor()
        PromoCode.objects.create(code='TEST-A1', sponsor=self.sp, status='used',
                                 used_at=timezone.now())
        PromoCode.objects.create(code='TEST-B2', sponsor=self.sp, status='active',
                                 created_at=timezone.now())

    def test_report_page(self):
        r = self.client.get('/admin/report/codes/')
        self.assertEqual(r.status_code, 200)
        self.assertContains(r, 'TEST-A1')

    def test_download_all(self):
        r = self.client.get('/admin/report/codes/?period=all&download=1')
        self.assertEqual(r.status_code, 200)
        self.assertEqual(r['Content-Disposition'].split(';')[0], 'attachment')
        content = r.content.decode('utf-8')
        self.assertIn('عدد الأكواد: 2', content)
        self.assertIn('المستخدمة: 1', content)

    def test_malformed_parameters_stay_200(self):
        bad_params = [
            {'period': 'month', 'month': 'abc', 'download': '1'},
            {'period': 'day', 'day': 'zzz', 'download': '1'},
            {'period': 'year', 'year': 'xx', 'download': '1'},
            {'period': 'nope', 'download': '1'},
        ]
        for params in bad_params:
            r = self.client.get('/admin/report/codes/', params)
            self.assertEqual(r.status_code, 200, msg=params)


class PageSmokeTests(TestCase):
    def test_public_pages(self):
        for path in ['/', '/mix/', '/search/', '/services/', '/sitemap.xml']:
            with self.subTest(path=path):
                self.assertEqual(self.client.get(path).status_code, 200)

    def test_dealers_page_hidden_by_default(self):
        settings = SiteSettings.load()
        settings.show_dealers_card = False
        settings.save()
        cache.delete(SITE_SETTINGS_CACHE_KEY)
        response = self.client.get('/dealers/')
        self.assertEqual(response.status_code, 302)
        self.assertEqual(response['Location'], '/')

    def test_dealers_page_filters_parts_region(self):
        settings = SiteSettings.load()
        settings.show_dealers_card = True
        settings.save()
        Dealer.objects.create(name='وكيل زيوت', dealer_type='oil', phone='07700000000', is_active=True)
        Dealer.objects.create(name='قطع ياباني', dealer_type='parts', parts_region='japanese', is_active=True)
        Dealer.objects.create(name='قطع ألماني', dealer_type='parts', parts_region='german', is_active=True)
        response = self.client.get('/dealers/', {'category': 'parts', 'parts_region': 'japanese'})
        self.assertContains(response, 'قطع ياباني')
        self.assertNotContains(response, 'قطع ألماني')

    def test_dealers_card_can_be_hidden_from_home(self):
        settings_obj = SiteSettings.load()
        settings_obj.show_dealers_card = False
        settings_obj.save()
        response = self.client.get('/')
        self.assertNotContains(response, 'وكلاء الزيوت وقطع الغيار')

    def test_security_headers_are_present(self):
        response = self.client.get('/')
        self.assertEqual(response['X-Frame-Options'], 'DENY')
        self.assertIn("default-src 'self'", response['Content-Security-Policy'])
        self.assertIn('camera=()', response['Permissions-Policy'])
        self.assertEqual(response['X-Permitted-Cross-Domain-Policies'], 'none')

    def test_admin_pages(self):
        self.client.force_login(User.objects.create_superuser('boss2', 'b2@example.com', 'pw'))
        for path in ['/admin/', '/admin/cars/promocode/', '/admin/report/codes/']:
            with self.subTest(path=path):
                self.assertEqual(self.client.get(path).status_code, 200)

    def test_cars_report_page(self):
        self.client.force_login(User.objects.create_superuser('boss3', 'b3@example.com', 'pw'))
        r = self.client.get('/admin/report/cars/')
        self.assertEqual(r.status_code, 200)
        self.assertContains(r, 'تقرير السيارات')

    def test_search_card_shows_engine_code_and_spark(self):
        car = _make_car(1, 'تويوتا', 'كورولا')
        car.engine_code = '2ZR-FE'
        car.spark = 'NGK Iridium'
        car.save()

        response = self.client.get('/search/', {'brand': 'Toyota'})
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, '2ZR-FE')
        self.assertContains(response, 'NGK Iridium')

    def test_car_export_excel_by_brand(self):
        self.client.force_login(User.objects.create_superuser('boss4', 'b4@example.com', 'pw'))
        _make_car(1, 'تويوتا', 'كورولا')
        _make_car(2, 'هيونداي', 'النترا')

        page = self.client.get('/admin/cars/carspecification/export-excel/')
        self.assertEqual(page.status_code, 200)
        self.assertContains(page, 'تصدير قاعدة بيانات السيارات')

        response = self.client.get('/admin/cars/carspecification/export-excel/', {'brand': 'تويوتا', 'download': '1'})
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response['Content-Disposition'].split(';')[0], 'attachment')
        self.assertEqual(response['Content-Type'], 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')

        from openpyxl import load_workbook
        sheet = load_workbook(BytesIO(response.content)).active
        rows = list(sheet.iter_rows(values_only=True))
        self.assertEqual(rows[0][0:5], ('Brand_EN', 'Brand_AR', 'Model_EN', 'Model_AR', 'Year'))
        self.assertIn('Engine Code', rows[0])
        self.assertIn('Spark', rows[0])
        self.assertEqual(rows[0][-1], 'id')
        self.assertEqual(len(rows), 2)
        self.assertEqual(rows[1][1], 'تويوتا')
