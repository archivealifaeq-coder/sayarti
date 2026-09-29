"""Shared, conservative Excel helpers for Django admin data tables."""

from __future__ import annotations

from datetime import datetime
from io import BytesIO

from django.core.exceptions import ValidationError
from django.http import HttpResponse
from django.utils import timezone
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter


EXCEL_CONTENT_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
MAX_EXCEL_FILE_SIZE = 10 * 1024 * 1024
MAX_EXCEL_ROWS = 20_000


def safe_excel_value(value):
    """Prevent spreadsheet formula injection in files downloaded by admins."""
    if isinstance(value, datetime) and timezone.is_aware(value):
        return timezone.localtime(value).replace(tzinfo=None)
    if isinstance(value, str) and value[:1] in ("=", "+", "-", "@"):
        return "'" + value
    return value


def excel_response(filename, sheets):
    """Build a styled XLSX response.

    ``sheets`` is an iterable of ``(title, headers, rows)``. Rows can be lists,
    tuples, or dictionaries whose keys match the headers.
    """
    workbook = Workbook()
    workbook.remove(workbook.active)

    for title, headers, rows in sheets:
        worksheet = workbook.create_sheet(title=str(title)[:31] or "Data")
        worksheet.sheet_view.rightToLeft = True
        worksheet.freeze_panes = "A2"
        worksheet.append(list(headers))
        for cell in worksheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="1D4ED8")

        for row in rows:
            values = [row.get(header, "") for header in headers] if isinstance(row, dict) else row
            worksheet.append([safe_excel_value(value) for value in values])

        worksheet.auto_filter.ref = worksheet.dimensions
        for column_index, header in enumerate(headers, start=1):
            values = [header]
            for cells in worksheet.iter_rows(min_row=2, min_col=column_index, max_col=column_index):
                values.append(cells[0].value)
            width = min(max(len(str(value or "")) for value in values) + 3, 48)
            worksheet.column_dimensions[get_column_letter(column_index)].width = max(width, 12)

    output = BytesIO()
    workbook.save(output)
    response = HttpResponse(output.getvalue(), content_type=EXCEL_CONTENT_TYPE)
    response["Content-Disposition"] = f'attachment; filename="{filename}"'
    response["X-Content-Type-Options"] = "nosniff"
    return response


def read_excel_rows(uploaded_file, preferred_sheets=()):
    """Read an XLSX upload as dictionaries without evaluating formulas."""
    filename = (getattr(uploaded_file, "name", "") or "").lower()
    if not filename.endswith(".xlsx"):
        raise ValidationError("يجب أن يكون الملف بصيغة .xlsx")
    if getattr(uploaded_file, "size", 0) > MAX_EXCEL_FILE_SIZE:
        raise ValidationError("حجم الملف أكبر من الحد المسموح (10 MB).")

    try:
        workbook = load_workbook(uploaded_file, read_only=True, data_only=True)
    except Exception as exc:
        raise ValidationError("تعذر قراءة ملف Excel. تأكد أن الملف غير تالف.") from exc

    sheet_lookup = {name.lower(): name for name in workbook.sheetnames}
    selected_name = next(
        (sheet_lookup[name.lower()] for name in preferred_sheets if name.lower() in sheet_lookup),
        workbook.sheetnames[0] if workbook.sheetnames else None,
    )
    if not selected_name:
        raise ValidationError("ملف Excel لا يحتوي على أوراق.")

    worksheet = workbook[selected_name]
    iterator = worksheet.iter_rows(values_only=True)
    headers = [str(value).strip() if value is not None else "" for value in next(iterator, ())]
    if not any(headers):
        raise ValidationError("صف العناوين في ملف Excel فارغ.")
    if len(headers) != len(set(headers)):
        raise ValidationError("ملف Excel يحتوي على عناوين أعمدة مكررة.")

    rows = []
    for row_number, values in enumerate(iterator, start=2):
        if row_number > MAX_EXCEL_ROWS + 1:
            raise ValidationError(f"عدد الصفوف يتجاوز الحد المسموح ({MAX_EXCEL_ROWS:,}).")
        row = {header: value for header, value in zip(headers, values) if header}
        if any(value not in (None, "") for value in row.values()):
            row["__row_number__"] = row_number
            rows.append(row)
    return rows


def text_value(row, name, default=""):
    value = row.get(name)
    if value is None:
        return default
    return str(value).strip()


def int_value(row, name, default=0):
    value = row.get(name)
    if value in (None, ""):
        return default
    try:
        return int(float(str(value).replace(",", "")))
    except (TypeError, ValueError) as exc:
        raise ValidationError(f"القيمة في العمود {name} يجب أن تكون رقماً.") from exc


def bool_value(row, name, default=True):
    value = row.get(name)
    if value in (None, ""):
        return default
    if isinstance(value, bool):
        return value
    normalized = str(value).strip().lower()
    if normalized in {"1", "true", "yes", "y", "نعم", "فعال", "مفعل"}:
        return True
    if normalized in {"0", "false", "no", "n", "لا", "متوقف", "غير مفعل"}:
        return False
    raise ValidationError(f"القيمة في العمود {name} يجب أن تكون نعم/لا.")


def datetime_value(row, name):
    value = row.get(name)
    if value in (None, ""):
        return None
    if isinstance(value, datetime):
        parsed = value
    else:
        try:
            parsed = datetime.fromisoformat(str(value).strip())
        except ValueError as exc:
            raise ValidationError(f"التاريخ في العمود {name} غير صحيح.") from exc
    if timezone.is_naive(parsed):
        parsed = timezone.make_aware(parsed, timezone.get_current_timezone())
    return parsed
